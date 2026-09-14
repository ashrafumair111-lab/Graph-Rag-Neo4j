"""Document loading — reads PDF / TXT / MD / DOCX files into LangChain Documents.

Supported extensions: .pdf, .txt, .md, .docx
"""
from __future__ import annotations

from pathlib import Path
from typing import List

from langchain_core.documents import Document

from config import get_logger

logger = get_logger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".md", ".docx"}


def _load_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n\n".join(pages)


def _load_docx(path: Path) -> str:
    from docx import Document as DocxDocument

    doc = DocxDocument(str(path))
    paragraphs = [p.text for p in doc.paragraphs if p.text]
    return "\n\n".join(paragraphs)


def _load_plain(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


def load_single(path: Path) -> Document:
    """Load one file into a single LangChain Document."""
    ext = path.suffix.lower()
    if ext == ".pdf":
        text = _load_pdf(path)
    elif ext == ".docx":
        text = _load_docx(path)
    elif ext in {".txt", ".md"}:
        text = _load_plain(path)
    else:
        raise ValueError(f"Unsupported file type: {ext}")

    if not text.strip():
        logger.warning("File '%s' produced no extractable text.", path)

    return Document(page_content=text, metadata={"source": str(path)})


def load_documents(path: str) -> List[Document]:
    """Load every supported file from a file path or a folder (recursively)."""
    root = Path(path)

    if root.is_file():
        files: List[Path] = [root]
    elif root.is_dir():
        files = sorted(
            p
            for p in root.rglob("*")
            if p.suffix.lower() in SUPPORTED_EXTENSIONS and not p.name.startswith("~$")
        )
    else:
        raise FileNotFoundError(f"Path does not exist: {path}")

    documents: List[Document] = []
    for file_path in files:
        try:
            documents.append(load_single(file_path))
            logger.info("Loaded: %s", file_path)
        except Exception as exc:  # noqa: BLE001 - keep the pipeline running
            logger.error("Failed to load %s: %s", file_path, exc)

    logger.info("Loaded %d document(s) from %s", len(documents), path)
    return documents